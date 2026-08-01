from __future__ import annotations

import io
import base64
import zipfile
from pathlib import Path

import pytest
from docx import Document

from autocite_mcp.document_ir import (
    classify_document_mode,
    locate_citations,
    parse_docx_ir,
    parse_markdown_ir,
    parse_pdf_ir,
    parse_text_ir,
)
from autocite_mcp.documents import DocumentLoadError, load_document_bytes
from autocite_mcp.engine import CitationEngine


FIXTURES = Path(__file__).parent / "fixtures" / "documents"


def _docx_with_footnote() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Argument. See 42 USC §1983")
    run = paragraph.add_run()
    marker = run._r.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}footnoteReference",
        {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id": "2"},
    )
    run._r.append(marker)
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Authority"
    stream = io.BytesIO()
    document.save(stream)
    source = zipfile.ZipFile(io.BytesIO(stream.getvalue()))
    document_xml = source.read("word/document.xml")
    footnotes = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:id="2"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t>Smith v. Jones, 123 F.3d 456 (9th Cir. 2020); 42 U.S.C. § 1983; 17 C.F.R. § 240.10b-5.</w:t></w:r></w:p></w:footnote>
</w:footnotes>""".encode()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            if item.filename == "word/document.xml":
                target.writestr(item, document_xml)
            else:
                target.writestr(item, source.read(item.filename))
        target.writestr("word/footnotes.xml", footnotes)
    return output.getvalue()


def test_plain_text_ir_has_stable_blocks_offsets_and_compatibility_mapping():
    text = "IN THE DISTRICT COURT\n\nArgument. See 42 USC §1983.\n\nConclusion."
    ir = parse_text_ir(text, filename="brief.txt")
    assert ir.to_text() == text
    paragraphs = ir.blocks_of_kind("paragraph")
    assert [block.block_id for block in paragraphs] == [
        "body:p:0",
        "body:p:1",
        "body:p:2",
    ]
    argument = paragraphs[1]
    assert ir.to_text()[argument.absolute_start : argument.absolute_end] == argument.text
    assert ir.block_at(argument.absolute_start + 2).block_id == argument.block_id
    assert ir.metadata.filename == "brief.txt"
    sentence = next(
        block for block in ir.blocks_of_kind("sentence") if block.parent_id == argument.block_id
    )
    assert sentence.block_local_start == argument.text.index("Argument")


def test_plain_text_preserves_explicit_numbered_footnote_syntax():
    text = "A proposition.[1]\n\n[1] See 42 U.S.C. § 1983."
    ir = parse_text_ir(text)
    assert ir.blocks_of_kind("footnote")[0].note_id == "1"
    assert ir.blocks_of_kind("footnote_reference")[0].note_id == "1"
    assert ir.to_text() == text


def test_plain_text_ir_sets_note_number_for_markdown_footnote_definitions():
    # review_document always parses through parse_text_ir (never
    # parse_markdown_ir), so markdown-style "[^1]:" footnote definitions must
    # get a note_number here too, or supra-note resolution (which matches on
    # location.note_number) can never succeed for these documents.
    text = (
        "Claim.[^1]\nLater, Author, supra note 1.\n\n"
        "[^1]: Jane Author, First Article, 12 Example L. Rev. 100 (2020)."
    )
    ir = parse_text_ir(text)
    notes = ir.blocks_of_kind("footnote")
    assert len(notes) == 1
    assert notes[0].note_id == "1"
    assert notes[0].note_number == "1"
    assert notes[0].text.startswith("Jane Author")
    assert ir.to_text() == text


def test_markdown_preserves_headings_links_block_quotes_and_explicit_notes():
    markdown = """# Seminar Paper

The claim follows from 42 USC §1983.[^source]

> A quoted proposition.

See [the archive](https://example.org/report).

[^source]: Introductory prose; Smith v. Jones, 123 F.3d 456 (9th Cir. 2020); 17 CFR §240.10b-5.
"""
    ir = parse_markdown_ir(markdown, filename="paper.md")
    assert len(ir.blocks_of_kind("heading")) == 1
    assert len(ir.blocks_of_kind("block_quotation")) == 1
    assert len(ir.blocks_of_kind("hyperlink")) == 1
    assert len(ir.blocks_of_kind("footnote_reference")) == 1
    notes = ir.blocks_of_kind("footnote")
    assert len(notes) == 1
    assert notes[0].note_id == "source"
    assert notes[0].metadata["markdown_identifier"] == "source"
    assert notes[0].parent_id is None
    assert "Smith v. Jones" in notes[0].text


def test_docx_ooxml_footnotes_are_separate_and_citations_keep_note_location():
    ir = parse_docx_ir(_docx_with_footnote(), filename="brief.docx")
    notes = ir.blocks_of_kind("footnote")
    assert len(notes) == 1
    assert notes[0].note_id == "2"
    assert notes[0].note_number == "1"
    assert notes[0].metadata["paragraph_order"] == 0
    references = ir.blocks_of_kind("footnote_reference")
    assert references[0].note_id == "2"
    assert references[0].metadata["target_block_id"] == notes[0].block_id
    assert len(ir.blocks_of_kind("table")) == 1

    citations = locate_citations(ir, CitationEngine())
    note_citations = [item for item in citations if item.location.note_id == "2"]
    assert len(note_citations) >= 2
    assert all(item.location.block_id == notes[0].block_id for item in note_citations)
    assert all(item.location.block_local_start >= 0 for item in note_citations)


def test_citation_opening_a_new_paragraph_keeps_its_block_location():
    # eyecite's case-name backward scan can absorb the leading "\n\n"
    # paragraph separator into the citation's own start, landing it in the
    # inter-block gap rather than inside the paragraph block that actually
    # contains it. Previously this made block_at() fail to find any
    # containing block, so the whole location (block_id, note_id,
    # reconstruction_confidence, etc.) was silently discarded even though
    # the citation is unambiguously inside the second paragraph.
    text = "Argument.\n\nSmith v. Jones, 123 F.3d 456 (9th Cir. 2020)."
    ir = parse_text_ir(text)
    citations = locate_citations(ir, CitationEngine())
    case_citation = next(item for item in citations if item.source_type == "case")
    assert case_citation.location.block_id is not None
    assert case_citation.location.reconstruction_confidence == "certain"
    assert case_citation.location.provenance == "parsed_document_structure"
    assert case_citation.location.block_local_start is not None
    assert case_citation.location.block_local_start >= 0


def test_pdf_preserves_pages_coordinates_and_uncertain_footnote_reconstruction(monkeypatch):
    class Box:
        height = 800

    class Page:
        media_box = Box()

        def __init__(self, number):
            self.number = number

        def extract_text(self, visitor_text=None):
            if visitor_text:
                visitor_text("Main proposition. 42 USC §1983.\n", None, [1, 0, 0, 1, 72, 700], None, 12)
                visitor_text("1 Smith v. Jones, 123 F.3d 456 (9th Cir. 2020).", None, [1, 0, 0, 1, 72, 45], None, 8)
            return "Main proposition. 42 USC §1983.\n1 Smith v. Jones, 123 F.3d 456 (9th Cir. 2020)."

    class Reader:
        def __init__(self, stream):
            self.pages = [Page(1), Page(2)]

    monkeypatch.setattr("autocite_mcp.document_ir.PdfReader", Reader)
    ir = parse_pdf_ir(b"%PDF-fake", filename="article.pdf")
    pages = ir.blocks_of_kind("page")
    assert [page.page_number for page in pages] == [1, 2]
    assert all(page.coordinates for page in pages)
    candidates = [block for block in ir.blocks if block.metadata.get("likely_footnote")]
    assert candidates
    assert all(block.reconstruction_confidence in {"low", "medium"} for block in candidates)
    assert any("uncertain" in warning.lower() for warning in ir.warnings)


def test_document_mode_classifier_reports_evidence_conflicts_and_confirmation():
    court = parse_text_ir(
        "IN THE UNITED STATES DISTRICT COURT\nPlaintiff moves for relief. See 42 U.S.C. § 1983.",
        filename="motion.txt",
    )
    selected = classify_document_mode(court, explicit_mode="whitepages")
    assert selected["selected_mode"] == "whitepages"
    assert selected["confidence"] == "high"
    assert "explicit_mode:whitepages" in selected["evidence"]
    assert selected["conflicting_evidence"]

    ambiguous = parse_text_ir("Research discussion with one citation, 42 U.S.C. § 1983.")
    result = classify_document_mode(ambiguous)
    assert result["selected_mode"] in {"bluepages", "whitepages"}
    assert result["user_confirmation_recommended"] is True


def test_academic_article_mentioning_supreme_court_classifies_as_whitepages():
    # Regression: an academic article discussing "the Supreme Court" plus a
    # couple of inline case citations used to false-positive into bluepages
    # because "supreme court" matched the court-filing regex and there were
    # no footnotes to offset it. An author byline is now also treated as
    # positive whitepages evidence.
    text = (
        "The Eroding Fourth Amendment\n"
        "By Greta Shope\n\n"
        "Introduction\n\n"
        "This article examines how the Supreme Court has narrowed Fourth "
        "Amendment protections over the past two decades. In Smith v. Jones, "
        "123 F.3d 456 (9th Cir. 2020), the court signaled a retreat from "
        "Katz v. United States, 389 U.S. 347 (1967). This trend is troubling "
        "for civil liberties scholars.\n\n"
        "This article proceeds in three parts. Part I surveys the doctrine. "
        "Part II analyzes recent Supreme Court decisions. Part III proposes "
        "reform.\n"
    )
    ir = parse_text_ir(text, filename="fourth_amendment_article.txt")
    result = classify_document_mode(ir)
    assert result["selected_mode"] == "whitepages"
    assert "author_byline" in result["evidence"]
    assert "court_filing_language" not in result["evidence"]


@pytest.mark.parametrize(
    "signature_block",
    [
        "Respectfully submitted,\n\nBy: John Smith\nAttorney for Plaintiff",
        "By: /s/ Jane Doe",
        "By Jane M. Doe, Esq.",
    ],
)
def test_legal_signature_blocks_do_not_trigger_the_byline_heuristic(signature_block):
    text = f"IN THE UNITED STATES DISTRICT COURT\n\nPlaintiff moves for relief.\n\n{signature_block}"
    ir = parse_text_ir(text, filename="motion.txt")
    result = classify_document_mode(ir)
    assert "author_byline" not in result["evidence"]
    assert result["selected_mode"] == "bluepages"


def test_load_document_bytes_keeps_flat_text_compatibility_and_exposes_ir():
    loaded = load_document_bytes("See 42 USC §1983.".encode(), "memo.txt")
    assert loaded.text == "See 42 USC §1983."
    assert loaded.ir.to_text() == loaded.text


@pytest.mark.parametrize(
    ("filename", "expected_mode"),
    [
        ("court_brief.txt", "bluepages"),
        ("legal_memorandum.txt", "bluepages"),
        ("law_review_article.md", "whitepages"),
        ("seminar_paper.md", "whitepages"),
        ("mixed_body_and_notes.md", "whitepages"),
    ],
)
def test_document_fixtures_preserve_citations_and_classify_mode(filename, expected_mode):
    loaded = load_document_bytes((FIXTURES / filename).read_bytes(), filename)
    result = classify_document_mode(loaded.ir)
    assert result["selected_mode"] == expected_mode
    citations = locate_citations(loaded.ir, CitationEngine())
    assert citations
    assert all(citation.location.block_id for citation in citations)


@pytest.mark.asyncio
async def test_uploaded_review_returns_structured_locations_without_breaking_flat_inventory():
    from autocite_mcp.tools import review_uploaded_document

    payload = _docx_with_footnote()
    result = await review_uploaded_document(
        {
            "file_name": "brief.docx",
            "data_base64": base64.b64encode(payload).decode(),
        },
        apply_safe_fixes=False,
    )
    assert result["citation_inventory"]
    assert result["structured_citation_inventory"]
    assert any(
        item["location"]["note_id"] == "2"
        for item in result["structured_citation_inventory"]
    )
    assert result["input_document"]["document_ir"]["footnote_count"] == 1
    # citation_count was a dead field, always 0: DocumentIR.citations defaults
    # to () and nothing ever called with_citations() before summary().
    assert result["document_ir"]["citation_count"] == len(result["structured_citation_inventory"])
    assert result["document_ir"]["citation_count"] > 0
    assert result["input_document"]["document_ir"]["citation_count"] == result["document_ir"]["citation_count"]


def test_scanned_pdf_still_returns_ocr_required(monkeypatch):
    class Page:
        def extract_text(self, visitor_text=None):
            return ""

    class Reader:
        def __init__(self, stream):
            self.pages = [Page()]

    monkeypatch.setattr("autocite_mcp.document_ir.PdfReader", Reader)
    with pytest.raises(DocumentLoadError) as exc:
        parse_pdf_ir(b"%PDF-image", filename="scan.pdf")
    assert exc.value.code == "ocr_required"


def test_pdf_page_count_is_capped(monkeypatch):
    # A crafted PDF can pack tens of thousands of tiny pages under the byte
    # cap; the page-count guard must reject it before per-page work begins.
    from autocite_mcp.document_ir import MAX_PDF_PAGES

    class Page:
        def extract_text(self, visitor_text=None):
            return "x"

    class Reader:
        def __init__(self, stream):
            self.pages = [Page() for _ in range(MAX_PDF_PAGES + 1)]

    monkeypatch.setattr("autocite_mcp.document_ir.PdfReader", Reader)
    with pytest.raises(DocumentLoadError) as exc:
        parse_pdf_ir(b"%PDF-huge", filename="huge.pdf")
    assert exc.value.code == "document_too_large"


def test_pdf_fragment_offsets_index_text_across_pages(monkeypatch):
    # Guards the O(1) running-offset rewrite: every reconstructed fragment
    # block's recorded span must still slice its own text out of ir.text.
    class Box:
        height = 800

    class Page:
        media_box = Box()

        def __init__(self, number):
            self.number = number

        def extract_text(self, visitor_text=None):
            if visitor_text:
                visitor_text(f"Page {self.number} body. 42 USC 1983.\n", None, [1, 0, 0, 1, 72, 700], None, 12)
                visitor_text(f"{self.number} A footnote fragment.", None, [1, 0, 0, 1, 72, 45], None, 8)
            return f"Page {self.number} body. 42 USC 1983.\n{self.number} A footnote fragment."

    class Reader:
        def __init__(self, stream):
            self.pages = [Page(i) for i in range(1, 6)]

    monkeypatch.setattr("autocite_mcp.document_ir.PdfReader", Reader)
    ir = parse_pdf_ir(b"%PDF-multi", filename="multi.pdf")
    for block in ir.blocks:
        if block.kind == "paragraph":
            assert ir.text[block.absolute_start : block.absolute_end] == block.text
