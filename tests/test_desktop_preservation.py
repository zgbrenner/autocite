from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.parts.hdrftr import HeaderPart

from autocite_mcp.desktop_preservation import (
    PreservationDesktopReviewController,
    PreservationDesktopReviewState,
    run_preservation_self_test,
)
from autocite_mcp.review_session import ReviewDecision


async def test_preservation_controller_retains_original_docx_structure(tmp_path: Path):
    source = tmp_path / "brief.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Privileged Draft"
    document.add_paragraph("See 42 USC §1983.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Record citation"
    table.cell(0, 1).text = "ECF 12"
    document.save(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    controller = PreservationDesktopReviewController()
    state = await controller.review_file(source, mode="bluepages")
    assert state.error_code is None
    assert state.result is not None
    assert state.source_bytes is not None

    destination = tmp_path / "review.docx"
    metadata = controller.export_docx(state, destination)

    assert metadata["preservation_mode"] == "original_docx"
    assert metadata["path"] == str(destination)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
    reviewed = Document(destination)
    assert reviewed.sections[0].header.paragraphs[0].text == "Privileged Draft"
    assert reviewed.tables[0].cell(0, 0).text == "Record citation"
    assert reviewed.tables[0].cell(0, 1).text == "ECF 12"


async def test_changed_docx_discards_review_and_requires_retry(
    tmp_path: Path, monkeypatch
):
    source = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("See 42 USC §1983.")
    document.save(source)
    original = source.read_bytes()
    real_read_bytes = Path.read_bytes
    source_reads = 0

    def changing_read_bytes(path: Path) -> bytes:
        nonlocal source_reads
        if path == source:
            source_reads += 1
            return original if source_reads == 1 else original + b"changed"
        return real_read_bytes(path)

    monkeypatch.setattr(Path, "read_bytes", changing_read_bytes)

    state = await PreservationDesktopReviewController().review_file(
        source, mode="bluepages"
    )

    assert source_reads == 2
    assert state.error_code == "docx_preservation_unavailable"
    assert state.result is None
    assert state.review_session is None
    assert state.source_bytes is None
    assert "review was discarded" in (state.error_message or "").lower()
    assert "review the current file again" in (state.error_message or "").lower()


async def test_desktop_export_applies_current_review_decisions(tmp_path: Path):
    source = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("See 42 USC §1983.")
    document.save(source)

    controller = PreservationDesktopReviewController()
    state = await controller.review_file(source, mode="bluepages")
    assert state.review_session is not None
    accepted = [
        item
        for item in state.review_session.items
        if item.decision is ReviewDecision.ACCEPTED
    ]
    assert accepted
    rejected_id = accepted[0].item_id

    destination = tmp_path / "review.docx"
    metadata = controller.export_docx(
        state,
        destination,
        decisions={rejected_id: ReviewDecision.REJECTED},
    )

    exported_items = metadata["review_session"]["items"]
    assert any(
        item["item_id"] == rejected_id and item["decision"] == "rejected"
        for item in exported_items
    )
    assert metadata["applied_edit_count"] == len(accepted) - 1


async def test_audit_report_uses_current_review_decisions(tmp_path: Path):
    source = tmp_path / "brief.txt"
    source.write_text("See 42 USC §1983.", encoding="utf-8")
    controller = PreservationDesktopReviewController()
    state = await controller.review_file(source, mode="bluepages")
    assert state.review_session is not None
    initial_accepted = [
        item
        for item in state.review_session.items
        if item.decision is ReviewDecision.ACCEPTED
    ]
    assert initial_accepted
    item_id = initial_accepted[0].item_id

    destination = tmp_path / "review.json"
    controller.export_json_report(
        state,
        destination,
        decisions={item_id: ReviewDecision.REJECTED},
    )
    report = json.loads(destination.read_text(encoding="utf-8"))

    assert any(
        item["item_id"] == item_id and item["decision"] == "rejected"
        for item in report["review_session"]["items"]
    )
    assert report["decision_summary"]["rejected"] == 1
    assert report["decision_summary"]["accepted_text_edits"] == len(initial_accepted) - 1
    assert report["summary"]["applied_edits"] == len(initial_accepted) - 1
    assert len(report["applied_edits"]) == len(initial_accepted) - 1
    assert all(edit["item_id"] != item_id for edit in report["export_plan"]["text_edits"])
    assert any(
        annotation["item_id"] == item_id
        and annotation["decision"] == "rejected"
        for annotation in report["export_plan"]["annotations"]
    )
    assert report["summary"]["mechanical_review_complete"] is False


async def test_non_docx_desktop_export_remains_a_labeled_fallback(tmp_path: Path):
    source = tmp_path / "brief.txt"
    source.write_text("See 42 USC §1983.", encoding="utf-8")

    controller = PreservationDesktopReviewController()
    state = await controller.review_file(source, mode="bluepages")
    destination = tmp_path / "review.docx"
    metadata = controller.export_docx(state, destination)

    assert destination.is_file()
    assert metadata["preservation_mode"] == "reconstructed_text"


async def test_preservation_self_test_exercises_existing_header_without_template(
    monkeypatch,
):
    def refuse_header_synthesis(*_args, **_kwargs):
        raise AssertionError("self-test must not synthesize a header from package templates")

    monkeypatch.setattr(HeaderPart, "new", refuse_header_synthesis)

    result = await run_preservation_self_test()

    assert result["status"] == "ok"
    assert result["preservation_mode"] == "original_docx"
    assert result["source_unchanged"] is True
    assert result["header_preserved"] is True
    assert result["table_preserved"] is True
    assert result["export_size_bytes"] > 0


def test_preserved_state_repr_does_not_expose_source_bytes():
    state = PreservationDesktopReviewState(
        original_text="original",
        corrected_text="corrected",
        result={},
        source_bytes=b"confidential document bytes",
    )

    assert "confidential document bytes" not in repr(state)
