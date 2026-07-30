from __future__ import annotations

import io
import json
import stat
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from autocite_mcp.docx_preservation import (
    DocxMappingError,
    DocxValidationError,
    apply_docx_export_plan,
    index_docx_text,
    validate_docx_package,
)
from autocite_mcp.review_session import (
    ExportPlan,
    PlannedAnnotation,
    PlannedTextEdit,
    ReviewDecision,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _document_bytes(*, split_target: bool = False) -> bytes:
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Privileged Draft"
    section.footer.paragraphs[0].text = "Confidential"

    paragraph = document.add_paragraph()
    paragraph.add_run("See 42 ")
    if split_target:
        paragraph.add_run("U")
        paragraph.add_run("SC")
    else:
        paragraph.add_run("USC")
    paragraph.add_run(" §1983 and Smith v. Jones.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Record citation"
    table.cell(0, 1).text = "ECF 12"
    table.cell(1, 0).text = "Exhibit A"
    table.cell(1, 1).text = "Page 4"

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _parts(payload: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _package(
    parts: dict[str, bytes],
    *,
    extra_entries: tuple[tuple[zipfile.ZipInfo | str, bytes], ...] = (),
) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
        for name_or_info, data in extra_entries:
            archive.writestr(name_or_info, data)
    return output.getvalue()


def _with_target_run_attribute(payload: bytes) -> bytes:
    parts = _parts(payload)
    root = etree.fromstring(parts["word/document.xml"])
    targets = root.xpath(
        ".//w:t[text()='USC']",
        namespaces={"w": W_NS},
    )
    assert len(targets) == 1
    run = targets[0].getparent()
    assert run is not None
    run.set(f"{{{W_NS}}}rsidR", "00ABCDEF")
    parts["word/document.xml"] = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
    )
    return _package(parts)


def _edit(index_text: str) -> PlannedTextEdit:
    start = index_text.index("USC")
    return PlannedTextEdit(
        item_id="edit-1",
        code="STATUTE_CODE_ABBREVIATION",
        start=start,
        end=start + 3,
        original="USC",
        replacement="U.S.C.",
        provenance="deterministic_logic",
    )


def _annotation(index_text: str) -> PlannedAnnotation:
    start = index_text.index("Smith")
    return PlannedAnnotation(
        item_id="annotation-1",
        code="PROPOSITION_PINCITE_REVIEW",
        start=start,
        end=start + len("Smith"),
        original="Smith",
        suggestion=None,
        message="Review whether this proposition requires a pinpoint citation.",
        severity="warning",
        confidence="medium",
        correction_level="review_required",
        provenance="deterministic_logic",
        rule="B10",
        missing_facts=(),
        decision=ReviewDecision.PENDING,
    )


def test_index_matches_autocite_body_and_table_text_order():
    index = index_docx_text(_document_bytes())

    assert index.text == (
        "See 42 USC §1983 and Smith v. Jones."
        "\n\nRecord citation\tECF 12\nExhibit A\tPage 4"
    )
    assert any(
        location.part_name == "word/document.xml"
        for location in index.locations
    )
    assert all(
        "header" not in location.part_name
        for location in index.locations
    )


def test_apply_plan_tracks_edit_and_preserves_unmodified_parts():
    payload = _document_bytes()
    before = _parts(payload)
    index = index_docx_text(payload)

    result = apply_docx_export_plan(
        payload,
        ExportPlan((_edit(index.text),), ()),
        tracked=True,
    )
    after = _parts(result.payload)

    assert result.applied_edit_ids == ("edit-1",)
    assert "word/document.xml" in result.modified_parts
    assert before["word/header1.xml"] == after["word/header1.xml"]
    assert before["word/footer1.xml"] == after["word/footer1.xml"]
    assert before["word/styles.xml"] == after["word/styles.xml"]
    document_xml = after["word/document.xml"].decode("utf-8")
    assert "<w:del" in document_xml
    assert "<w:ins" in document_xml
    assert "U.S.C." in document_xml
    assert validate_docx_package(result.payload).valid is True


def test_tracked_edit_preserves_run_level_attributes_on_new_runs():
    payload = _with_target_run_attribute(_document_bytes())
    index = index_docx_text(payload)

    result = apply_docx_export_plan(
        payload,
        ExportPlan((_edit(index.text),), ()),
        tracked=True,
    )

    document_xml = _parts(result.payload)["word/document.xml"]
    assert document_xml.count(b'rsidR="00ABCDEF"') == 2


def test_untracked_edit_replaces_text_without_revision_markup():
    payload = _document_bytes()
    index = index_docx_text(payload)

    result = apply_docx_export_plan(
        payload,
        ExportPlan((_edit(index.text),), ()),
        tracked=False,
    )

    document_xml = _parts(result.payload)["word/document.xml"].decode("utf-8")
    assert "U.S.C." in document_xml
    assert "<w:del" not in document_xml
    assert "<w:ins" not in document_xml


def test_cross_run_edit_is_refused_instead_of_corrupting_document():
    payload = _document_bytes(split_target=True)
    index = index_docx_text(payload)

    with pytest.raises(DocxMappingError, match="unambiguous"):
        apply_docx_export_plan(
            payload,
            ExportPlan((_edit(index.text),), ()),
            tracked=True,
        )


def test_review_annotation_creates_word_comment_for_reliable_body_range():
    payload = _document_bytes()
    index = index_docx_text(payload)

    result = apply_docx_export_plan(
        payload,
        ExportPlan((), (_annotation(index.text),)),
        tracked=True,
    )
    parts = _parts(result.payload)

    assert result.anchored_annotation_ids == ("annotation-1",)
    assert result.unanchored_annotations == ()
    assert "word/comments.xml" in parts
    comments = parts["word/comments.xml"].decode("utf-8")
    assert "PROPOSITION_PINCITE_REVIEW" in comments
    assert "pinpoint citation" in comments
    document_xml = parts["word/document.xml"].decode("utf-8")
    assert "commentRangeStart" in document_xml
    assert "commentRangeEnd" in document_xml
    assert "commentReference" in document_xml


def test_annotation_that_crosses_runs_remains_in_audit_instead_of_forcing_markup():
    payload = _document_bytes(split_target=True)
    index = index_docx_text(payload)
    start = index.text.index("USC") - 1
    annotation = PlannedAnnotation(
        item_id="annotation-cross-run",
        code="MANUAL_REVIEW",
        start=start,
        end=start + 5,
        original=index.text[start : start + 5],
        suggestion=None,
        message="Review this range.",
        severity="warning",
        confidence="medium",
        correction_level="review_required",
        provenance="deterministic_logic",
        rule="",
        missing_facts=(),
        decision=ReviewDecision.PENDING,
    )

    result = apply_docx_export_plan(
        payload,
        ExportPlan((), (annotation,)),
        tracked=True,
    )

    assert result.anchored_annotation_ids == ()
    assert [
        item.item_id for item in result.unanchored_annotations
    ] == ["annotation-cross-run"]
    assert "word/comments.xml" not in _parts(result.payload)


def test_validator_rejects_missing_internal_relationship_target():
    payload = _document_bytes()
    parts = _parts(payload)
    rels_name = "word/_rels/document.xml.rels"
    rels = parts[rels_name].replace(
        b"</Relationships>",
        (
            b'<Relationship Id="rIdBroken" Type="urn:test" '
            b'Target="missing.xml"/></Relationships>'
        ),
    )
    output = _package(
        {
            name: rels if name == rels_name else data
            for name, data in parts.items()
        }
    )

    with pytest.raises(
        DocxValidationError,
        match="missing relationship target",
    ):
        validate_docx_package(output, raise_on_error=True)


def test_validator_rejects_path_traversal_package_member():
    payload = _package(
        _parts(_document_bytes()),
        extra_entries=(("../outside.txt", b"forbidden"),),
    )

    with pytest.raises(DocxValidationError, match="unsafe package path"):
        validate_docx_package(payload, raise_on_error=True)


def test_validator_rejects_symbolic_link_package_member():
    link = zipfile.ZipInfo("word/linked.xml")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    payload = _package(
        _parts(_document_bytes()),
        extra_entries=((link, b"document.xml"),),
    )

    with pytest.raises(DocxValidationError, match="symbolic-link"):
        validate_docx_package(payload, raise_on_error=True)


def test_validator_rejects_xml_document_type_and_entity_declarations():
    parts = _parts(_document_bytes())
    parts["word/document.xml"] = (
        b'<!DOCTYPE document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
        + parts["word/document.xml"]
    )

    with pytest.raises(
        DocxValidationError,
        match="prohibited document type declaration",
    ):
        validate_docx_package(
            _package(parts),
            raise_on_error=True,
        )


def test_validator_rejects_duplicate_comment_ids():
    parts = _parts(_document_bytes())
    parts["word/comments.xml"] = (
        f'<w:comments xmlns:w="{W_NS}">'
        '<w:comment w:id="0"/><w:comment w:id="0"/>'
        "</w:comments>"
    ).encode("utf-8")

    with pytest.raises(DocxValidationError, match="duplicate comment IDs"):
        validate_docx_package(
            _package(parts),
            raise_on_error=True,
        )


def test_result_metadata_is_json_serializable():
    payload = _document_bytes()
    index = index_docx_text(payload)
    result = apply_docx_export_plan(
        payload,
        ExportPlan(
            (_edit(index.text),),
            (_annotation(index.text),),
        ),
        tracked=True,
    )

    serialized = json.dumps(result.as_dict(), sort_keys=True)
    assert "original_docx" in serialized
    assert "edit-1" in serialized
    assert "annotation-1" in serialized
