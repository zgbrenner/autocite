from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

from .desktop import (
    DesktopReviewController,
    DesktopReviewState,
    _atomic_write,
    build_desktop_report,
)
from .desktop_export import build_desktop_docx_export
from .review_session import ReviewDecision, ReviewItemKind, ReviewSession
from .tools import export_review_docx


@dataclass(frozen=True)
class PreservationDesktopReviewState(DesktopReviewState):
    source_bytes: bytes | None = field(default=None, repr=False, compare=False)
    review_session: ReviewSession | None = field(default=None, repr=False, compare=False)
    preservation_warning: str | None = None


class PreservationDesktopReviewController(DesktopReviewController):
    """Desktop controller that preserves the original Word package when possible."""

    @staticmethod
    def _upgrade_state(
        state: DesktopReviewState,
        *,
        source_bytes: bytes | None = None,
        preservation_warning: str | None = None,
    ) -> PreservationDesktopReviewState:
        session = ReviewSession.from_result(state.result) if state.result is not None else None
        return PreservationDesktopReviewState(
            original_text=state.original_text,
            corrected_text=state.corrected_text,
            result=state.result,
            error_code=state.error_code,
            error_message=state.error_message,
            source_path=state.source_path,
            source_bytes=source_bytes,
            review_session=session,
            preservation_warning=preservation_warning,
        )

    @staticmethod
    def _preservation_failure(
        state: DesktopReviewState,
        message: str,
    ) -> PreservationDesktopReviewState:
        return PreservationDesktopReviewState(
            original_text="",
            corrected_text="",
            result=None,
            error_code="docx_preservation_unavailable",
            error_message=message,
            source_path=state.source_path,
            preservation_warning=message,
        )

    async def review_file(
        self,
        path: Path,
        *,
        document_type: str = "auto",
        mode: str = "auto",
        jurisdiction: str | None = None,
        use_local_model: bool = False,
    ) -> PreservationDesktopReviewState:
        state = await super().review_file(
            path,
            document_type=document_type,
            mode=mode,
            jurisdiction=jurisdiction,
            use_local_model=use_local_model,
        )
        if state.result is None or state.source_path is None:
            return self._upgrade_state(state)
        input_document = state.result.get("input_document")
        source_format = (
            str(input_document.get("source_format") or "")
            if isinstance(input_document, dict)
            else ""
        )
        if source_format != "docx":
            return self._upgrade_state(state)
        try:
            source_bytes = await asyncio.to_thread(state.source_path.read_bytes)
        except OSError as exc:
            return self._preservation_failure(
                state,
                "AutoCite completed citation analysis but could not retain the original "
                "Word package safely. The review was discarded and no export is available. "
                f"Close other programs using the file, then review it again. Technical detail: {exc}",
            )
        expected_hash = (
            str(input_document.get("sha256") or "")
            if isinstance(input_document, dict)
            else ""
        )
        actual_hash = hashlib.sha256(source_bytes).hexdigest()
        if not expected_hash or actual_hash != expected_hash:
            return self._preservation_failure(
                state,
                "The Word document changed while AutoCite was reviewing it. The review was "
                "discarded and no export is available. Review the current file again.",
            )
        return self._upgrade_state(state, source_bytes=source_bytes)

    @staticmethod
    def _fallback_docx(original: str, corrected: str, tracked: bool) -> bytes:
        artifact = export_review_docx(
            original,
            corrected,
            tracked=tracked,
            filename="autocite-review.docx",
        )
        return base64.b64decode(str(artifact["data_base64"]), validate=True)

    @staticmethod
    def _session_with_decisions(
        state: DesktopReviewState,
        decisions: Mapping[str, ReviewDecision | str] | None,
    ) -> ReviewSession:
        if isinstance(state, PreservationDesktopReviewState) and state.review_session is not None:
            session = state.review_session
        elif state.result is not None:
            session = ReviewSession.from_result(state.result)
        else:
            raise ValueError("a completed review is required before export")
        return session.with_decisions(decisions) if decisions else session

    def export_docx(
        self,
        state: DesktopReviewState,
        destination: Path,
        *,
        decisions: Mapping[str, ReviewDecision | str] | None = None,
    ) -> dict[str, Any]:
        result = self._require_result(state)
        target = self._validate_destination(state, destination, ".docx")
        source_bytes = (
            state.source_bytes
            if isinstance(state, PreservationDesktopReviewState)
            else None
        )
        session = self._session_with_decisions(state, decisions)
        export = build_desktop_docx_export(
            result=result,
            source_bytes=source_bytes,
            fallback_builder=self._fallback_docx,
            tracked=True,
            decisions=session.decisions(),
        )
        _atomic_write(target, export.payload)
        return {
            **export.metadata,
            "filename": target.name,
            "mime_type": (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            "path": str(target),
            "size_bytes": len(export.payload),
        }

    def export_json_report(
        self,
        state: DesktopReviewState,
        destination: Path,
        *,
        decisions: Mapping[str, ReviewDecision | str] | None = None,
    ) -> dict[str, Any]:
        self._require_result(state)
        target = self._validate_destination(state, destination, ".json")
        report = build_desktop_report(state)
        session = self._session_with_decisions(state, decisions)
        plan = session.export_plan()
        accepted_edit_items = [
            item.as_dict()
            for item in session.items
            if item.kind is ReviewItemKind.TEXT_EDIT
            and item.decision is ReviewDecision.ACCEPTED
            and item.suggestion is not None
        ]

        accepted = sum(
            item.decision is ReviewDecision.ACCEPTED for item in session.items
        )
        rejected = sum(
            item.decision is ReviewDecision.REJECTED for item in session.items
        )
        pending = sum(
            item.decision is ReviewDecision.PENDING for item in session.items
        )
        unsupported = sum(
            item.correction_level == "unsupported" for item in session.items
        )
        review_items = sum(
            item.kind is ReviewItemKind.ANNOTATION
            and (
                item.decision is not ReviewDecision.ACCEPTED
                or item.correction_level == "unsupported"
            )
            for item in session.items
        )

        summary = report.get("summary")
        summary = dict(summary) if isinstance(summary, dict) else {}
        summary.update(
            {
                "applied_edits": len(plan.text_edits),
                "review_items": review_items,
                "unsupported_items": unsupported,
                "remaining_mechanical_issues": len(plan.annotations),
                "mechanical_review_complete": not plan.annotations,
            }
        )
        report["summary"] = summary
        report["applied_edits"] = accepted_edit_items
        report["review_session"] = session.as_dict()
        report["export_plan"] = plan.as_dict()
        report["decision_summary"] = {
            "total": len(session.items),
            "accepted": accepted,
            "rejected": rejected,
            "pending": pending,
            "accepted_text_edits": len(plan.text_edits),
            "unresolved_export_annotations": len(plan.annotations),
            "unsupported": unsupported,
        }
        if isinstance(state, PreservationDesktopReviewState):
            report["docx_preservation"] = {
                "ready": state.source_bytes is not None,
                "warning": state.preservation_warning,
            }
        payload = (
            json.dumps(
                report,
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
                default=str,
            )
            + "\n"
        ).encode("utf-8")
        _atomic_write(target, payload)
        return {"path": str(target), "size_bytes": len(payload)}


async def run_preservation_self_test() -> dict[str, Any]:
    """Exercise review, original-DOCX mutation, package validation, and fidelity."""

    from docx import Document

    with tempfile.TemporaryDirectory(prefix="autocite-preservation-self-test-") as directory:
        root = Path(directory)
        source = root / "preservation-self-test.docx"
        document = Document()
        document.sections[0].header.paragraphs[0].text = "AutoCite Self-Test"
        document.add_paragraph("See 42 USC §1983.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Structure"
        table.cell(0, 1).text = "Preserved"
        document.save(source)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        controller = PreservationDesktopReviewController()
        state = await controller.review_file(source, mode="bluepages")
        if state.error_code or state.result is None:
            raise RuntimeError(state.error_message or "preservation self-test review failed")
        if state.source_bytes is None:
            raise RuntimeError(
                state.preservation_warning
                or "preservation self-test did not retain the original DOCX"
            )
        destination = root / "preservation-self-test-review.docx"
        metadata = controller.export_docx(state, destination)
        if metadata.get("preservation_mode") != "original_docx":
            raise RuntimeError("preservation self-test used a reconstruction fallback")
        if not destination.is_file() or destination.stat().st_size <= 0:
            raise RuntimeError("preservation self-test export was not created")

        reviewed = Document(destination)
        header_preserved = (
            reviewed.sections[0].header.paragraphs[0].text == "AutoCite Self-Test"
        )
        table_preserved = (
            bool(reviewed.tables)
            and reviewed.tables[0].cell(0, 0).text == "Structure"
            and reviewed.tables[0].cell(0, 1).text == "Preserved"
        )
        source_unchanged = hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
        if not header_preserved or not table_preserved or not source_unchanged:
            raise RuntimeError("preservation self-test detected a fidelity or source mutation failure")
        return {
            "status": "ok",
            "preservation_mode": "original_docx",
            "source_unchanged": source_unchanged,
            "header_preserved": header_preserved,
            "table_preserved": table_preserved,
            "applied_edit_count": int(metadata.get("applied_edit_count", 0)),
            "export_size_bytes": destination.stat().st_size,
        }
