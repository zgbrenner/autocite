from __future__ import annotations

import asyncio
import io
import ipaddress
import mimetypes
import socket
from dataclasses import dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import TYPE_CHECKING
from urllib.parse import urljoin, urlparse

import httpx
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

if TYPE_CHECKING:
    from .document_ir import DocumentIR

MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MIN_PDF_TEXT_CHARS = 20
MAX_DOWNLOAD_REDIRECTS = 4


class DocumentLoadError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class DocumentInput:
    text: str
    filename: str
    mime_type: str
    source_format: str
    warnings: tuple[str, ...] = ()
    ir: DocumentIR | None = None


def validate_download_url(url: str) -> str:
    """Reject non-HTTPS and obvious private-network file URLs."""
    parsed = urlparse(url)
    if parsed.scheme.lower() != "https":
        raise DocumentLoadError("unsafe_download_url", "Remote document URLs must use HTTPS.")
    if parsed.username or parsed.password:
        raise DocumentLoadError("unsafe_download_url", "Remote document URLs cannot contain credentials.")
    host = (parsed.hostname or "").strip().lower().rstrip(".")
    if not host:
        raise DocumentLoadError("unsafe_download_url", "Remote document URL has no hostname.")
    if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
        raise DocumentLoadError("unsafe_download_url", "Local and private-network hostnames are not allowed.")
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    if address and (
        address.is_private
        or address.is_loopback
        or address.is_link_local
        or address.is_multicast
        or address.is_reserved
        or address.is_unspecified
    ):
        raise DocumentLoadError("unsafe_download_url", "Private-network addresses are not allowed.")
    return url


async def _resolve_pinned_url(url: str) -> tuple[str, str]:
    """Resolve and validate the URL's host, returning an IP-pinned URL.

    The request is made against the exact validated address (with SNI and Host
    preserved), so a DNS-rebinding server cannot serve a public IP to the
    validation lookup and a private one to the connection.
    """
    parsed = urlparse(validate_download_url(url))
    host = parsed.hostname or ""
    port = parsed.port or 443
    try:
        results = await asyncio.to_thread(
            socket.getaddrinfo,
            host,
            port,
            type=socket.SOCK_STREAM,
        )
    except OSError as exc:
        raise DocumentLoadError("download_dns_error", "Remote document hostname could not be resolved.") from exc
    for result in results:
        address = ipaddress.ip_address(result[4][0])
        if (
            address.is_private
            or address.is_loopback
            or address.is_link_local
            or address.is_multicast
            or address.is_reserved
            or address.is_unspecified
        ):
            raise DocumentLoadError(
                "unsafe_download_url",
                "Remote document hostname resolves to a private or reserved network.",
            )
    pinned_ip = results[0][4][0]
    pinned_netloc = f"[{pinned_ip}]:{port}" if ":" in pinned_ip else f"{pinned_ip}:{port}"
    return parsed._replace(netloc=pinned_netloc).geturl(), host


async def download_document_url(
    url: str,
    *,
    timeout: float = 30.0,
    client_factory: type[httpx.AsyncClient] = httpx.AsyncClient,
) -> tuple[bytes, str | None]:
    """Download an authorized file URL with redirect, SSRF, and size safeguards."""
    current = validate_download_url(url)
    async with client_factory(timeout=timeout, follow_redirects=False) as client:
        for _ in range(MAX_DOWNLOAD_REDIRECTS + 1):
            pinned_url, host = await _resolve_pinned_url(current)
            async with client.stream(
                "GET",
                pinned_url,
                headers={"Accept": "*/*", "Host": host},
                extensions={"sni_hostname": host},
            ) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    location = response.headers.get("location")
                    if not location:
                        raise DocumentLoadError("download_redirect_error", "Remote document redirect has no location.")
                    current = validate_download_url(urljoin(current, location))
                    continue
                response.raise_for_status()
                length = response.headers.get("content-length")
                if length and int(length) > MAX_DOCUMENT_BYTES:
                    raise DocumentLoadError("document_too_large", "Remote document exceeds the 15 MB limit.")
                chunks: list[bytes] = []
                total = 0
                async for chunk in response.aiter_bytes():
                    total += len(chunk)
                    if total > MAX_DOCUMENT_BYTES:
                        raise DocumentLoadError("document_too_large", "Remote document exceeds the 15 MB limit.")
                    chunks.append(chunk)
                return b"".join(chunks), response.headers.get("content-type")
    raise DocumentLoadError("download_redirect_error", "Remote document exceeded the redirect limit.")


def _normalized_mime(filename: str, mime_type: str | None) -> str:
    if mime_type:
        return mime_type.split(";", 1)[0].strip().lower()
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def _docx_text(payload: bytes) -> str:
    document = Document(io.BytesIO(payload))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _pdf_text(payload: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(payload))
        parts = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        raise DocumentLoadError(
            "ocr_required",
            "The PDF could not be parsed as a text PDF. OCR or a searchable PDF is required.",
        ) from exc
    text = "\n\n".join(part for part in parts if part).strip()
    if len(text) < MIN_PDF_TEXT_CHARS:
        raise DocumentLoadError(
            "ocr_required",
            "The PDF contains too little extractable text and appears scanned or image-only.",
        )
    return text


def load_document_bytes(data: bytes, filename: str, mime_type: str | None = None) -> DocumentInput:
    """Load a supported document without persisting it."""
    if not data:
        raise DocumentLoadError("empty_document", "The uploaded document is empty.")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentLoadError(
            "document_too_large",
            f"Documents are limited to {MAX_DOCUMENT_BYTES // (1024 * 1024)} MB.",
        )
    safe_name = Path(filename or "document.txt").name
    mime = _normalized_mime(safe_name, mime_type)
    suffix = Path(safe_name).suffix.lower()
    from .document_ir import parse_docx_ir, parse_markdown_ir, parse_pdf_ir, parse_text_ir

    if mime in {"text/plain", "text/markdown"} or suffix in {".txt", ".md", ".markdown"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentLoadError("invalid_text_encoding", "Text documents must use UTF-8.") from exc
        source_format = "markdown" if suffix in {".md", ".markdown"} or mime == "text/markdown" else "text"
        ir = (
            parse_markdown_ir(text, filename=safe_name, mime_type=mime)
            if source_format == "markdown"
            else parse_text_ir(text, filename=safe_name, mime_type=mime)
        )
    elif (
        mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == ".docx"
    ):
        try:
            ir = parse_docx_ir(data, filename=safe_name)
            text = ir.to_text()
        except Exception as exc:
            raise DocumentLoadError("invalid_docx", "The DOCX file could not be read.") from exc
        source_format = "docx"
    elif mime == "application/pdf" or suffix == ".pdf":
        try:
            ir = parse_pdf_ir(data, filename=safe_name)
            text = ir.to_text()
        except DocumentLoadError:
            raise
        except Exception as exc:
            raise DocumentLoadError("invalid_pdf", "The PDF file could not be read.") from exc
        source_format = "pdf"
    else:
        raise DocumentLoadError(
            "unsupported_document_type",
            "Supported files are TXT, Markdown, DOCX, and text-based PDF.",
        )

    if not text.strip():
        raise DocumentLoadError("empty_document", "The document contains no readable text.")
    return DocumentInput(
        text=text,
        filename=safe_name,
        mime_type=mime,
        source_format=source_format,
        warnings=tuple(ir.warnings)
        + ("Input formatting is normalized to text for citation analysis; layout may not be preserved in exports.",),
        ir=ir,
    )


def _text_node(tag: str, value: str) -> OxmlElement:
    node = OxmlElement(tag)
    if value[:1].isspace() or value[-1:].isspace():
        node.set(qn("xml:space"), "preserve")
    node.text = value
    return node


def _revision(parent: OxmlElement, tag: str, value: str, revision_id: int) -> None:
    if not value:
        return
    change = OxmlElement(tag)
    change.set(qn("w:id"), str(revision_id))
    change.set(qn("w:author"), "AutoCite")
    change.set(qn("w:date"), datetime.now(timezone.utc).isoformat())
    run = OxmlElement("w:r")
    run.append(_text_node("w:delText" if tag == "w:del" else "w:t", value))
    change.append(run)
    parent.append(change)


def build_review_docx(original_text: str, corrected_text: str, *, tracked: bool = True) -> bytes:
    """Build a simple review DOCX; tracked mode emits Word insertion/deletion markup."""
    document = Document()
    paragraph = document.add_paragraph()
    if not tracked:
        paragraph.add_run(corrected_text)
    else:
        matcher = SequenceMatcher(None, original_text, corrected_text)
        revision_id = 1
        for opcode, i1, i2, j1, j2 in matcher.get_opcodes():
            if opcode == "equal":
                paragraph.add_run(original_text[i1:i2])
            elif opcode == "delete":
                _revision(paragraph._p, "w:del", original_text[i1:i2], revision_id)
                revision_id += 1
            elif opcode == "insert":
                _revision(paragraph._p, "w:ins", corrected_text[j1:j2], revision_id)
                revision_id += 1
            elif opcode == "replace":
                _revision(paragraph._p, "w:del", original_text[i1:i2], revision_id)
                revision_id += 1
                _revision(paragraph._p, "w:ins", corrected_text[j1:j2], revision_id)
                revision_id += 1
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
