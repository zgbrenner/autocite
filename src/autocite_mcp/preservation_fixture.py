from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree

_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_PACKAGE_RELATIONSHIPS_NS = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_WORDPROCESSING_NS = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)
_OFFICE_RELATIONSHIPS_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_HEADER_RELATIONSHIP_TYPE = f"{_OFFICE_RELATIONSHIPS_NS}/header"
_HEADER_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"
)
_HEADER_PART_NAME = "word/header1.xml"


def _serialize_xml(element: etree._Element) -> bytes:
    return etree.tostring(
        element,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def _next_relationship_id(relationships: etree._Element) -> str:
    existing = {
        relationship.get("Id")
        for relationship in relationships
        if relationship.get("Id")
    }
    candidate = "rIdAutoCiteSelfTestHeader"
    counter = 1
    while candidate in existing:
        counter += 1
        candidate = f"rIdAutoCiteSelfTestHeader{counter}"
    return candidate


def _attach_existing_header(path: Path, text: str) -> None:
    """Attach a real header part without asking python-docx to synthesize one."""

    with ZipFile(path, "r") as archive:
        members = archive.infolist()
        payloads = {member.filename: archive.read(member.filename) for member in members}

    required = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/_rels/document.xml.rels",
    }
    missing = sorted(required - payloads.keys())
    if missing:
        raise ValueError(f"preservation fixture is missing DOCX parts: {missing}")
    if _HEADER_PART_NAME in payloads:
        raise ValueError("preservation fixture already contains header1.xml")

    content_types = etree.fromstring(payloads["[Content_Types].xml"])
    override = content_types.xpath(
        "./ct:Override[@PartName='/word/header1.xml']",
        namespaces={"ct": _CONTENT_TYPES_NS},
    )
    if not override:
        header_override = etree.SubElement(
            content_types,
            etree.QName(_CONTENT_TYPES_NS, "Override"),
        )
        header_override.set("PartName", "/word/header1.xml")
        header_override.set("ContentType", _HEADER_CONTENT_TYPE)

    relationships = etree.fromstring(payloads["word/_rels/document.xml.rels"])
    relationship_id = _next_relationship_id(relationships)
    header_relationship = etree.SubElement(
        relationships,
        etree.QName(_PACKAGE_RELATIONSHIPS_NS, "Relationship"),
    )
    header_relationship.set("Id", relationship_id)
    header_relationship.set("Type", _HEADER_RELATIONSHIP_TYPE)
    header_relationship.set("Target", "header1.xml")

    document = etree.fromstring(payloads["word/document.xml"])
    body = document.find(etree.QName(_WORDPROCESSING_NS, "body"))
    if body is None:
        raise ValueError("preservation fixture has no Word document body")
    section_properties = body.find(etree.QName(_WORDPROCESSING_NS, "sectPr"))
    if section_properties is None:
        raise ValueError("preservation fixture has no section properties")
    header_reference = etree.Element(
        etree.QName(_WORDPROCESSING_NS, "headerReference")
    )
    header_reference.set(etree.QName(_WORDPROCESSING_NS, "type"), "default")
    header_reference.set(
        etree.QName(_OFFICE_RELATIONSHIPS_NS, "id"), relationship_id
    )
    section_properties.insert(0, header_reference)

    header = etree.Element(
        etree.QName(_WORDPROCESSING_NS, "hdr"),
        nsmap={"w": _WORDPROCESSING_NS},
    )
    paragraph = etree.SubElement(header, etree.QName(_WORDPROCESSING_NS, "p"))
    run = etree.SubElement(paragraph, etree.QName(_WORDPROCESSING_NS, "r"))
    text_element = etree.SubElement(run, etree.QName(_WORDPROCESSING_NS, "t"))
    text_element.text = text

    payloads["[Content_Types].xml"] = _serialize_xml(content_types)
    payloads["word/_rels/document.xml.rels"] = _serialize_xml(relationships)
    payloads["word/document.xml"] = _serialize_xml(document)

    temporary = path.with_name(f".{path.name}.autocite-fixture.tmp")
    try:
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as archive:
            for member in members:
                archive.writestr(member, payloads[member.filename])
            archive.writestr(
                _HEADER_PART_NAME,
                _serialize_xml(header),
                compress_type=ZIP_DEFLATED,
            )
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def create_preservation_fixture(path: Path) -> None:
    """Create a DOCX containing a citation, table, and pre-existing header."""

    document = Document()
    document.add_paragraph("See 42 USC §1983.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Structure"
    table.cell(0, 1).text = "Preserved"
    document.save(path)
    _attach_existing_header(path, "AutoCite Self-Test")
