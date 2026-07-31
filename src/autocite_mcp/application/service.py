from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from .context_reducer import ContextReducer, DeterministicLegalContextReducer
from .engine_adapter import AutoCiteReviewEngine, ReviewEngine
from .models import (
    DecisionState,
    DocumentSession,
    ReviewDecision,
    ReviewJob,
    ReviewJobState,
    SessionStatus,
)
from .store import RevisionConflict, SessionStore


class ApplicationServiceError(RuntimeError):
    pass


class UnsupportedDocumentFormat(ApplicationServiceError):
    pass


class InvalidReviewDecision(ApplicationServiceError):
    pass


@dataclass(frozen=True, slots=True)
class ExportArtifact:
    export_id: str
    format: str
    filename: str
    media_type: str
    content: bytes
    sha256: str
    path: str | None = None

    def summary(self) -> dict[str, Any]:
        return {
            "export_id": self.export_id,
            "format": self.format,
            "filename": self.filename,
            "media_type": self.media_type,
            "size_bytes": len(self.content),
            "sha256": self.sha256,
            "path": self.path,
        }


class ApplicationService:
    """Canonical stateful service consumed by MCP, HTTP, and Tauri clients."""

    def __init__(
        self,
        store: SessionStore | None = None,
        *,
        engine: ReviewEngine | None = None,
        reducer: ContextReducer | None = None,
    ) -> None:
        self.store = store or SessionStore()
        self.engine = engine or AutoCiteReviewEngine()
        self.reducer = reducer or DeterministicLegalContextReducer()

    def create_session(
        self,
        *,
        title: str,
        text: str,
        source_format: str = "md",
        metadata: dict[str, Any] | None = None,
    ) -> DocumentSession:
        return self.store.create_session(
            title=title,
            source_format=source_format,
            original_text=text,
            metadata=metadata,
        )

    def import_path(self, path: str | Path) -> DocumentSession:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        suffix = source.suffix.lower()
        if suffix in {".txt", ".md", ".markdown"}:
            text = source.read_text(encoding="utf-8-sig")
            source_format = "md" if suffix in {".md", ".markdown"} else "txt"
        elif suffix == ".docx":
            text = self._docx_to_markdown(source)
            source_format = "docx"
        elif suffix == ".pdf":
            text = self._pdf_to_markdown(source)
            source_format = "pdf"
        else:
            raise UnsupportedDocumentFormat(
                f"Unsupported document format: {source.suffix or '(none)'}"
            )
        return self.store.create_session(
            title=source.stem,
            source_format=source_format,
            original_text=text,
            metadata={
                "source_path": str(source),
                "canonical_format": "markdown",
                "importer": "application_service",
            },
        )

    @staticmethod
    def _docx_to_markdown(path: Path) -> str:
        document = Document(path)
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.rstrip()
            if not text:
                blocks.append("")
                continue
            style_name = (paragraph.style.name if paragraph.style else "").lower()
            heading_match = re.search(r"heading\s+(\d+)", style_name)
            if heading_match:
                level = min(max(int(heading_match.group(1)), 1), 6)
                blocks.append(f"{'#' * level} {text}")
            elif "list bullet" in style_name:
                blocks.append(f"- {text}")
            elif "list number" in style_name:
                blocks.append(f"1. {text}")
            else:
                blocks.append(text)
        for table in document.tables:
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized = [row + [""] * (width - len(row)) for row in rows]
            blocks.append("| " + " | ".join(normalized[0]) + " |")
            blocks.append("| " + " | ".join(["---"] * width) + " |")
            blocks.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
        return "\n\n".join(block for block in blocks).strip()

    @staticmethod
    def _pdf_to_markdown(path: Path) -> str:
        reader = PdfReader(path)
        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            pages.append(f"<!-- page: {index} -->\n\n{text}")
        if not pages:
            raise UnsupportedDocumentFormat(
                "The PDF contains no extractable text and requires OCR."
            )
        return "\n\n".join(pages)

    def get_session(self, session_id: str) -> dict[str, Any]:
        session = self.store.get_session(session_id)
        decisions = [item.to_dict() for item in self.store.list_decisions(session_id)]
        payload = session.to_dict()
        payload["decisions"] = decisions
        payload["word_count"] = _word_count(session.working_text)
        payload["citation_count"] = _citation_count(session.working_text)
        return payload

    def list_sessions(self, *, limit: int = 50, offset: int = 0) -> dict[str, Any]:
        sessions = self.store.list_sessions(limit=limit, offset=offset)
        return {
            "sessions": [self._session_summary(session) for session in sessions],
            "limit": min(max(int(limit), 1), 200),
            "offset": max(int(offset), 0),
        }

    @staticmethod
    def _session_summary(session: DocumentSession) -> dict[str, Any]:
        return {
            "id": session.id,
            "title": session.title,
            "source_format": session.source_format,
            "revision": session.revision,
            "status": session.status.value,
            "word_count": _word_count(session.working_text),
            "citation_count": _citation_count(session.working_text),
            "created_at": session.created_at,
            "updated_at": session.updated_at,
        }

    def update_session(
        self,
        session_id: str,
        *,
        expected_revision: int,
        text: str | None = None,
        title: str | None = None,
        editor_state: dict[str, Any] | None = None,
    ) -> DocumentSession:
        metadata_patch = {"editor_state": editor_state} if editor_state is not None else None
        return self.store.update_session(
            session_id,
            expected_revision=expected_revision,
            working_text=text,
            title=title,
            metadata_patch=metadata_patch,
            clear_review=text is not None,
        )

    def context_preview(
        self,
        session_id: str,
        *,
        focus_spans: list[tuple[int, int]] | None = None,
        max_characters: int | None = None,
        include_text: bool = False,
    ) -> dict[str, Any]:
        session = self.store.get_session(session_id)
        result = self.reducer.reduce(
            session.working_text,
            focus_spans=focus_spans,
            max_characters=max_characters,
        )
        payload = result.to_dict()
        if not include_text:
            payload.pop("text", None)
        payload["session_id"] = session.id
        payload["session_revision"] = session.revision
        return payload

    async def review_session(
        self,
        session_id: str,
        *,
        expected_revision: int,
        mode: str | None = None,
        jurisdiction: str | None = None,
        deep_review: bool = False,
        use_local_model: bool = False,
    ) -> ReviewJob:
        session = self.store.get_session(session_id)
        if session.revision != expected_revision:
            raise RevisionConflict(session_id, expected_revision, session.revision)
        job = self.store.create_review_job(session_id, session.revision)
        self.store.set_status(
            session_id, SessionStatus.REVIEWING, expected_revision=session.revision
        )
        self.store.update_review_job(
            job.id,
            state=ReviewJobState.RUNNING,
            progress=0.05,
        )
        try:
            # The deterministic engine always receives the complete document.
            # Reduction metadata is prepared for bounded SLM and retrieval tasks,
            # never as a lossy replacement for citation extraction.
            reduction = self.reducer.reduce(session.working_text)
            raw_review = await self.engine.review(
                session.working_text,
                mode=mode,
                jurisdiction=jurisdiction,
                deep_review=deep_review,
                use_local_model=use_local_model,
            )
            normalized_issues = _normalize_issues(raw_review, session.working_text)
            review = dict(raw_review)
            review["application_issues"] = normalized_issues
            review["session_revision"] = session.revision
            review["context_reduction"] = {
                "strategy": reduction.metadata.get("strategy"),
                "metrics": reduction.metrics.to_dict(),
                "selected_ranges": [
                    {"start": item.start, "end": item.end, "reason": item.reason}
                    for item in reduction.selected_ranges
                ],
                "reversible": True,
                "used_for_deterministic_review": False,
            }
            self.store.save_review(
                session_id,
                expected_revision=session.revision,
                review=review,
            )
            summary = {
                "session_id": session.id,
                "session_revision": session.revision,
                "issue_count": len(normalized_issues),
                "safe_auto_fix_count": sum(
                    1 for issue in normalized_issues if issue["safe_to_apply"]
                ),
                "mode": raw_review.get("mode", mode),
                "jurisdiction": raw_review.get("jurisdiction", jurisdiction),
            }
            return self.store.update_review_job(
                job.id,
                state=ReviewJobState.COMPLETED,
                progress=1.0,
                result_summary=summary,
            )
        except Exception as exc:
            self.store.set_status(session_id, SessionStatus.ERROR)
            self.store.update_review_job(
                job.id,
                state=ReviewJobState.FAILED,
                progress=1.0,
                error=f"{type(exc).__name__}: {exc}",
            )
            raise

    def review_session_sync(self, *args: Any, **kwargs: Any) -> ReviewJob:
        try:
            asyncio.get_running_loop()
        except RuntimeError:
            return asyncio.run(self.review_session(*args, **kwargs))
        raise RuntimeError("review_session_sync cannot run inside an event loop")

    def get_review_job(self, job_id: str) -> dict[str, Any]:
        return self.store.get_review_job(job_id).to_dict()

    def decide_issue(
        self,
        session_id: str,
        *,
        issue_id: str,
        decision: DecisionState | str,
        expected_revision: int,
        rationale: str | None = None,
    ) -> dict[str, Any]:
        state = DecisionState(decision)
        session = self.store.get_session(session_id)
        if session.revision != expected_revision:
            raise RevisionConflict(session_id, expected_revision, session.revision)
        issue = _find_application_issue(session.latest_review, issue_id)
        if state is DecisionState.REJECTED:
            saved = self.store.save_decision(
                ReviewDecision(
                    session_id=session_id,
                    issue_id=issue_id,
                    state=state,
                    expected_revision=session.revision,
                    replacement=issue.get("replacement"),
                    rationale=rationale,
                )
            )
            return {
                "decision": saved.to_dict(),
                "session": self._session_summary(session),
                "review_stale": False,
            }
        if state is DecisionState.PENDING:
            saved = self.store.save_decision(
                ReviewDecision(
                    session_id=session_id,
                    issue_id=issue_id,
                    state=state,
                    expected_revision=session.revision,
                    rationale=rationale,
                )
            )
            return {
                "decision": saved.to_dict(),
                "session": self._session_summary(session),
                "review_stale": False,
            }
        return self.apply_safe_issues(
            session_id,
            issue_ids=[issue_id],
            expected_revision=expected_revision,
            rationale=rationale,
        )

    def apply_safe_issues(
        self,
        session_id: str,
        *,
        issue_ids: list[str],
        expected_revision: int,
        rationale: str | None = None,
    ) -> dict[str, Any]:
        session = self.store.get_session(session_id)
        if session.revision != expected_revision:
            raise RevisionConflict(session_id, expected_revision, session.revision)
        if not session.latest_review:
            raise InvalidReviewDecision("The session has no current review")
        issues = [_find_application_issue(session.latest_review, issue_id) for issue_id in issue_ids]
        _validate_non_overlapping_issues(issues)
        text = session.working_text
        for issue in sorted(issues, key=lambda item: int(item["start"]), reverse=True):
            if not issue.get("safe_to_apply"):
                raise InvalidReviewDecision(
                    f"Issue {issue['id']} is not eligible for automatic application"
                )
            start = int(issue["start"])
            end = int(issue["end"])
            original = str(issue.get("original", ""))
            replacement = issue.get("replacement")
            if not isinstance(replacement, str):
                raise InvalidReviewDecision(f"Issue {issue['id']} has no replacement")
            if start < 0 or end < start or end > len(text):
                raise InvalidReviewDecision(f"Issue {issue['id']} has an invalid span")
            if text[start:end] != original:
                raise InvalidReviewDecision(
                    f"Issue {issue['id']} no longer matches the current document"
                )
            text = text[:start] + replacement + text[end:]

        updated = self.store.update_session(
            session_id,
            expected_revision=session.revision,
            working_text=text,
            status=SessionStatus.READY,
            clear_review=True,
        )
        decisions = []
        for issue in issues:
            saved = self.store.save_decision(
                ReviewDecision(
                    session_id=session_id,
                    issue_id=str(issue["id"]),
                    state=DecisionState.ACCEPTED,
                    expected_revision=updated.revision,
                    replacement=str(issue["replacement"]),
                    rationale=rationale,
                )
            )
            decisions.append(saved.to_dict())
        return {
            "decisions": decisions,
            "session": self._session_summary(updated),
            "review_stale": True,
            "requires_review_refresh": True,
        }

    def export_session(
        self,
        session_id: str,
        *,
        output_format: str,
        destination: str | Path | None = None,
    ) -> ExportArtifact:
        session = self.store.get_session(session_id)
        normalized = output_format.lower().lstrip(".")
        stem = _safe_filename(session.title) or "autocite-document"
        if normalized in {"txt", "text"}:
            extension = "txt"
            media_type = "text/plain; charset=utf-8"
            content = session.working_text.encode("utf-8")
        elif normalized in {"md", "markdown"}:
            extension = "md"
            media_type = "text/markdown; charset=utf-8"
            content = session.working_text.encode("utf-8")
        elif normalized == "docx":
            extension = "docx"
            media_type = (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            content = _markdown_to_docx(session.working_text)
        elif normalized == "pdf":
            extension = "pdf"
            media_type = "application/pdf"
            content = _plain_text_pdf(session.working_text, title=session.title)
        else:
            raise UnsupportedDocumentFormat(f"Unsupported export format: {output_format}")

        filename = f"{stem}-autocite.{extension}"
        target_path: str | None = None
        if destination is not None:
            target = Path(destination).expanduser().resolve()
            if target.exists() and target.is_dir():
                target = target / filename
            elif not target.suffix:
                target = target.with_suffix(f".{extension}")
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(content)
            target_path = str(target)
            filename = target.name
        digest = hashlib.sha256(content).hexdigest()
        export_id = self.store.record_export(
            session_id=session.id,
            session_revision=session.revision,
            output_format=extension,
            sha256=digest,
            path=target_path,
        )
        return ExportArtifact(
            export_id=export_id,
            format=extension,
            filename=filename,
            media_type=media_type,
            content=content,
            sha256=digest,
            path=target_path,
        )


def _normalize_issues(review: dict[str, Any], text: str) -> list[dict[str, Any]]:
    candidates: list[Any] = []
    for key in (
        "issues",
        "review_items",
        "findings",
        "deterministic_issues",
        "remaining_issues",
        "applied_edits",
    ):
        value = review.get(key)
        if isinstance(value, list):
            candidates.extend(value)
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(candidates):
        if not isinstance(raw, dict):
            continue
        start, end = _extract_span(raw)
        original = _first_string(
            raw,
            "original",
            "citation",
            "source_text",
            "matched_text",
            "old_text",
        )
        if start is None and original:
            start = text.find(original)
            if start < 0:
                start = None
        if start is not None and end is None and original is not None:
            end = start + len(original)
        if start is None or end is None or start < 0 or end > len(text) or end < start:
            continue
        if original is None:
            original = text[start:end]
        replacement = _first_string(
            raw,
            "replacement",
            "corrected",
            "corrected_text",
            "suggestion",
            "suggested_citation",
            "new_text",
        )
        issue_code = _first_string(raw, "issue_code", "code", "rule", "kind", "type")
        message = _first_string(raw, "message", "explanation", "description", "reason")
        level = _first_string(raw, "correction_level", "level", "confidence") or "review"
        safe = bool(raw.get("safe_to_apply") or raw.get("auto_fixable"))
        safe = safe or level.lower() in {"safe_auto_fix", "safe", "high_confidence"}
        safe = safe and replacement is not None and text[start:end] == original
        identity_material = json.dumps(
            [start, end, original, replacement, issue_code, index], ensure_ascii=False
        )
        issue_id = str(raw.get("id") or hashlib.sha256(identity_material.encode()).hexdigest()[:20])
        if issue_id in seen:
            continue
        seen.add(issue_id)
        normalized.append(
            {
                "id": issue_id,
                "start": start,
                "end": end,
                "original": original,
                "replacement": replacement,
                "issue_code": issue_code,
                "message": message,
                "correction_level": level,
                "safe_to_apply": safe,
                "provenance": raw.get("provenance", "deterministic_engine"),
                "raw": raw,
            }
        )
    normalized.sort(key=lambda item: (int(item["start"]), int(item["end"])))
    return normalized


def _extract_span(raw: dict[str, Any]) -> tuple[int | None, int | None]:
    start = raw.get("start", raw.get("start_offset"))
    end = raw.get("end", raw.get("end_offset"))
    span = raw.get("span")
    if isinstance(span, dict):
        start = span.get("start", span.get("start_offset", start))
        end = span.get("end", span.get("end_offset", end))
    elif isinstance(span, (list, tuple)) and len(span) >= 2:
        start, end = span[0], span[1]
    try:
        normalized_start = int(start) if start is not None else None
        normalized_end = int(end) if end is not None else None
    except (TypeError, ValueError):
        return None, None
    return normalized_start, normalized_end


def _first_string(mapping: dict[str, Any], *keys: str) -> str | None:
    for key in keys:
        value = mapping.get(key)
        if isinstance(value, str):
            return value
    return None


def _find_application_issue(
    review: dict[str, Any] | None, issue_id: str
) -> dict[str, Any]:
    if not review:
        raise InvalidReviewDecision("The session has no current review")
    issues = review.get("application_issues")
    if not isinstance(issues, list):
        raise InvalidReviewDecision("The current review has no application issues")
    for issue in issues:
        if isinstance(issue, dict) and str(issue.get("id")) == issue_id:
            return issue
    raise InvalidReviewDecision(f"Review issue not found: {issue_id}")


def _validate_non_overlapping_issues(issues: list[dict[str, Any]]) -> None:
    ordered = sorted(issues, key=lambda item: (int(item["start"]), int(item["end"])))
    for previous, current in zip(ordered, ordered[1:], strict=False):
        if int(current["start"]) < int(previous["end"]):
            raise InvalidReviewDecision(
                f"Issues {previous['id']} and {current['id']} overlap"
            )


def _word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text, flags=re.UNICODE))


def _citation_count(text: str) -> int:
    patterns = (
        r"\b\d{1,4}\s+(?:U\.S\.|F\.\s?(?:2d|3d|4th)|S\.\s*Ct\.)\s+\d+",
        r"\b\d+\s+(?:U\.S\.C\.|C\.F\.R\.)\s*§",
        r"\b(?:Id\.|supra\b|infra\b)",
    )
    return sum(len(re.findall(pattern, text, flags=re.IGNORECASE)) for pattern in patterns)


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "", value).strip().replace(" ", "-")
    return re.sub(r"-+", "-", cleaned)[:100]


def _markdown_to_docx(markdown: str) -> bytes:
    document = Document()
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif re.match(r"^[-*+]\s+", stripped):
            document.add_paragraph(re.sub(r"^[-*+]\s+", "", stripped), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", stripped):
            document.add_paragraph(
                re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number"
            )
        elif stripped.startswith("|") and stripped.endswith("|"):
            # Table reconstruction is intentionally conservative.  Keeping the
            # Markdown row visible is safer than guessing merged cells.
            document.add_paragraph(stripped)
        else:
            document.add_paragraph(stripped)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _plain_text_pdf(text: str, *, title: str) -> bytes:
    """Create a dependency-free, text-first PDF for portable releases.

    The rich editor can use the platform print pipeline for exact visual
    output.  This fallback keeps backend PDF export available offline and
    converts unsupported glyphs to readable WinAnsi approximations.
    """

    replacements = str.maketrans(
        {
            "“": '"',
            "”": '"',
            "‘": "'",
            "’": "'",
            "–": "-",
            "—": "-",
            "…": "...",
            "§": "Sec.",
        }
    )
    normalized = text.translate(replacements)
    lines: list[str] = []
    for paragraph in normalized.splitlines() or [""]:
        wrapped = textwrap.wrap(paragraph, width=88, replace_whitespace=False)
        lines.extend(wrapped or [""])
    page_capacity = 48
    pages = [lines[index : index + page_capacity] for index in range(0, len(lines), page_capacity)] or [[""]]

    objects: list[bytes] = []

    def add_object(payload: bytes) -> int:
        objects.append(payload)
        return len(objects)

    catalog_id = add_object(b"")
    pages_id = add_object(b"")
    font_id = add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    for page_number, page_lines in enumerate(pages, start=1):
        commands = ["BT", "/F1 11 Tf", "54 738 Td", "14 TL"]
        header = f"{title}  |  AutoCite export  |  Page {page_number}"
        for position, line in enumerate([header, "", *page_lines]):
            escaped = (
                line.encode("cp1252", errors="replace")
                .decode("latin-1")
                .replace("\\", "\\\\")
                .replace("(", "\\(")
                .replace(")", "\\)")
            )
            commands.append(f"({escaped}) Tj")
            if position < len(page_lines) + 1:
                commands.append("T*")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        stream_id = add_object(
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
        page_id = add_object(
            (
                f"<< /Type /Page /Parent {pages_id} 0 R "
                f"/MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {stream_id} 0 R >>"
            ).encode()
        )
        page_ids.append(page_id)

    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    objects[catalog_id - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, payload in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode())
        output.extend(payload)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)
